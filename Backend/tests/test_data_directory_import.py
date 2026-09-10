from openpyxl import Workbook, load_workbook

from hamamooz.apps.evaluations.models import MetricScore, MonthlyEvaluation
from hamamooz.apps.imports.data_directory import DataDirectoryScanner, selected_manifest_for_class
from hamamooz.apps.imports.direct_data_ingest import ingest_data_directory
from hamamooz.apps.imports.models import (
    ClassSourceSelection,
    DataSourceConflict,
    DataSourceManifest,
    DataSourceRow,
    SubjectExamResult,
)


def _save_workbook(path, *, class_code="701", include_classification=True):
    workbook = Workbook()
    class_sheet = workbook.active
    class_sheet.title = "کلاس‌بندی" if include_classification else "راهنما"
    if include_classification:
        class_sheet.append(["عنوان"])
        class_sheet.append([])
        class_sheet.append(["ردیف", "کد کلاس", "نام کلاس", "پایه تحصیلی"])
        class_sheet.append([1, class_code, "کلاس آزمون", "هفتم"])

    student_sheet = workbook.create_sheet("دانش‌آموزان")
    student_sheet.append(["فهرست"])
    student_sheet.append([])
    student_sheet.append(["ردیف", "کد ملی", "نام", "نام خانوادگی", "کد کلاس", "پایه"])
    student_sheet.append([1, "0123456789", "علی", "آزمون", class_code, "هفتم"])

    evaluation_sheet = workbook.create_sheet("ثبت اطلاعات")
    evaluation_sheet.append(["ارزیابی"])
    evaluation_sheet.append([])
    evaluation_sheet.append(["ردیف", "ماه", "کد ملی", "کد کلاس", "EDU_01 | نمرات درسی"])
    evaluation_sheet.append([1, "شهریور", "0123456789", class_code, 5])
    workbook.save(path)


def test_scanner_preserves_raw_rows_and_normalizes_monthly_provenance(base_data, tmp_path):
    _save_workbook(tmp_path / "701.xlsx")

    result = DataDirectoryScanner(base_data["school1"], tmp_path).scan()

    assert result["files_scanned"] == 1
    manifest = DataSourceManifest.objects.get(source_file="701.xlsx")
    assert manifest.status == DataSourceManifest.Status.VALID
    assert manifest.detected_classes == ["701"]
    evaluation = DataSourceRow.objects.get(
        manifest=manifest, row_kind=DataSourceRow.RowKind.EVALUATION
    )
    assert evaluation.source_file == "701.xlsx"
    assert evaluation.source_row == 4
    assert evaluation.national_id == "0123456789"
    assert evaluation.month_no == 3
    assert evaluation.normalized_data["month_title"] == "شهریور"
    assert evaluation.normalized_data["metrics"] == {"EDU_01": 5}
    assert evaluation.raw_values["values"][1] == "شهریور"


def test_scanner_keeps_formula_cells_in_raw_values(tmp_path, base_data):
    path = tmp_path / "formula.xlsx"
    _save_workbook(path)
    workbook = load_workbook(path)
    workbook["ثبت اطلاعات"]["E4"] = "=1+1"
    workbook.save(path)
    workbook.close()

    DataDirectoryScanner(base_data["school1"], tmp_path).scan()

    row = DataSourceRow.objects.get(
        source_file="formula.xlsx", sheet_name="ثبت اطلاعات", source_row=4
    )
    assert row.raw_values["values"][4] == "=1+1"
    # There is no cached result in a workbook created by openpyxl; the normalized
    # projection is allowed to remain empty while the formula evidence survives.
    assert row.normalized_data["metrics"]["EDU_01"] is None


def test_file_without_classification_is_incomplete_and_not_an_official_source(base_data, tmp_path):
    _save_workbook(tmp_path / "902.xlsx", class_code="902", include_classification=False)

    DataDirectoryScanner(base_data["school1"], tmp_path).scan()

    manifest = DataSourceManifest.objects.get(source_file="902.xlsx")
    assert manifest.status == DataSourceManifest.Status.INCOMPLETE
    assert manifest.errors[0]["code"] == "missing_class_definition"
    ClassSourceSelection.objects.create(
        school=base_data["school1"], class_code="902", manifest=manifest
    )
    assert selected_manifest_for_class(base_data["school1"], "902") is None


def test_direct_ingestor_writes_valid_comprehensive_rows_to_official_tables(base_data, tmp_path):
    workbook = Workbook()
    class_sheet = workbook.active
    class_sheet.title = "کلاس‌بندی"
    class_sheet.append(["عنوان"])
    class_sheet.append(["راهنما"])
    class_sheet.append([])
    class_sheet.append(
        ["ردیف", "کد مدرسه", "سال تحصیلی", "کد کلاس", "نام کلاس", "پایه تحصیلی", "ظرفیت"]
    )
    class_sheet.append(
        [1, base_data["school1"].code, base_data["year"].code, "7-c", "هفتم ج", "هفتم", 35]
    )

    student_sheet = workbook.create_sheet("دانش‌آموزان")
    student_sheet.append(["فهرست"])
    student_sheet.append(["راهنما"])
    student_sheet.append([])
    student_sheet.append(
        [
            "ردیف",
            "کد محلی",
            "کد ملی",
            "شماره دانش‌آموزی",
            "نام",
            "نام خانوادگی",
            "جنسیت",
            "تاریخ تولد",
            "کد کلاس",
        ]
    )
    student_sheet.append([1, 1, "0012345680", "103", "علی", "آزمون", "پسر", "1391/01/01", "7-c"])

    evaluation_sheet = workbook.create_sheet("ثبت اطلاعات")
    evaluation_sheet.append(["ثبت ارزیابی"])
    evaluation_sheet.append(["راهنما"])
    evaluation_sheet.append([])
    evaluation_sheet.append(
        ["ردیف", "ماه", "کد محلی", "کد ملی", "نام و نام خانوادگی", "کد کلاس", "EDU_01 | نمرات درسی"]
    )
    evaluation_sheet.append([1, "شهریور", 1, "0012345680", "علی آزمون", "7-c", 5])
    workbook.save(tmp_path / "direct.xlsx")
    workbook.close()

    result = ingest_data_directory(
        base_data["school1"],
        excel_directory=tmp_path,
        photo_directory=tmp_path / "photos",
        recorded_by=base_data["manager"],
    )

    assert result["phase_order"] == ["comprehensive", "subject_exam", "photos"]
    assert result["comprehensive_files"] == 1
    assert result["comprehensive"][0]["students_written"] == 1
    assert result["comprehensive"][0]["enrollments_written"] == 1
    assert result["comprehensive"][0]["metric_scores_written"] == 1
    enrollment = base_data["school1"].school_enrollments.get(
        student__national_id="0012345680", academic_year=base_data["year"]
    )
    assert MonthlyEvaluation.objects.filter(enrollment=enrollment, month_no=3).count() == 1
    assert (
        MetricScore.objects.filter(
            evaluation__enrollment=enrollment, metric_code="EDU_01", value=5
        ).count()
        == 1
    )


def test_overlapping_class_and_student_create_conflicts_without_implicit_selection(
    base_data, tmp_path
):
    _save_workbook(tmp_path / "first.xlsx")
    _save_workbook(tmp_path / "second.xlsx")

    DataDirectoryScanner(base_data["school1"], tmp_path).scan()

    conflicts = DataSourceConflict.objects.filter(school=base_data["school1"], status="open")
    assert conflicts.filter(
        conflict_type=DataSourceConflict.ConflictType.CLASS_OVERLAP, class_code="701"
    ).exists()
    assert conflicts.filter(
        conflict_type=DataSourceConflict.ConflictType.STUDENT_OVERLAP,
        class_code="701",
        national_id="0123456789",
    ).exists()
    assert selected_manifest_for_class(base_data["school1"], "701") is None

    selected = DataSourceManifest.objects.get(source_file="first.xlsx")
    ClassSourceSelection.objects.create(
        school=base_data["school1"], class_code="701", manifest=selected
    )
    assert selected_manifest_for_class(base_data["school1"], "701") == selected


def test_protected_data_source_endpoints_scan_list_conflicts_and_select(
    api_client, base_data, tmp_path, settings
):
    _save_workbook(tmp_path / "first.xlsx")
    _save_workbook(tmp_path / "second.xlsx")
    settings.DATA_EXCEL_DIRECTORY = str(tmp_path)
    headers = {"HTTP_X_SCHOOL_ID": str(base_data["school1"].id)}
    api_client.force_authenticate(base_data["manager"])

    scan = api_client.post("/api/v1/data-sources/scan/", {}, format="json", **headers)
    assert scan.status_code == 200
    assert scan.data["files_scanned"] == 2

    manifests = api_client.get("/api/v1/data-sources/", **headers)
    assert manifests.status_code == 200
    assert manifests.data["count"] == 2
    conflicts = api_client.get("/api/v1/data-source-conflicts/?class_code=701", **headers)
    assert conflicts.status_code == 200
    assert conflicts.data["count"] >= 2

    manifest = DataSourceManifest.objects.get(source_file="first.xlsx")
    select = api_client.post(
        "/api/v1/class-source-selections/",
        {"school": str(base_data["school1"].id), "class_code": "701", "manifest": str(manifest.id)},
        format="json",
        **headers,
    )
    assert select.status_code == 201, select.data
    assert select.data["selected_by"] == base_data["manager"].id
    assert selected_manifest_for_class(base_data["school1"], "701") == manifest

    api_client.force_authenticate(base_data["teacher1"])
    hidden = api_client.get("/api/v1/data-source-conflicts/", **headers)
    assert hidden.status_code == 200
    assert hidden.data["count"] == 0


def test_subject_exam_rows_have_a_scoped_read_only_audit_endpoint(api_client, base_data):
    SubjectExamResult.objects.create(
        organization=base_data["organization"],
        school=base_data["school1"],
        source_file="کارنامه تفصیلی آزمون تابستانه هفتم.xlsx",
        source_checksum="a" * 64,
        source_sheet="Sheet1",
        source_row=2,
        first_name="علی",
        last_name="آزمون",
        national_id_raw="0012345678",
        national_id="0012345678",
        grade_name="هفتم",
        grade_order=7,
        class_name="7-a",
        subject_name="ریاضی",
        status=SubjectExamResult.Status.UNMATCHED,
        raw_values={"نمره": "=16"},
        errors=[{"code": "student_not_found"}],
        error_count=1,
    )

    api_client.force_authenticate(base_data["manager"])
    response = api_client.get(
        "/api/v1/subject-exam-results/",
        {"source_file": "کارنامه تفصیلی آزمون تابستانه هفتم.xlsx"},
        HTTP_X_SCHOOL_ID=str(base_data["school1"].id),
    )

    assert response.status_code == 200, response.data
    assert response.data["count"] == 1
    assert response.data["results"][0]["status"] == SubjectExamResult.Status.UNMATCHED
    assert response.data["results"][0]["raw_values"]["نمره"] == "=16"
