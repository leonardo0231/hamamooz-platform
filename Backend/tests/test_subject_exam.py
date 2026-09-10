from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook

from hamamooz.apps.imports.models import DataSourceManifest, SubjectExamResult
from hamamooz.apps.imports.services.subject_exam import (
    SUBJECT_EXAM_FILENAMES,
    SubjectExamWorkbookParser,
    discover_subject_exam_workbooks,
    ingest_subject_exam_workbooks,
)
from hamamooz.apps.students.models import Student

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_EXCEL = PROJECT_ROOT / "Data" / "Excel"
SUBJECT_EXAM_ROW_COUNTS = {
    "کارنامه تفصیلی آزمون تابستانه نهم.xlsx": 1272,
    "کارنامه تفصیلی آزمون تابستانه هشتم.xlsx": 1288,
    "کارنامه تفصیلی آزمون تابستانه هفتم.xlsx": 1330,
}


def _file_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_minimal_subject_exam(path: Path, score: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(
        [
            "نام",
            "نام خانوادگی",
            "نام کاربری",
            "کلاس",
            "درس",
            "تعداد سوال",
            "صحیح",
            "غلط",
            "سفید",
            "درصد",
            "بالاترین درصد",
            "رتبه",
            "تراز",
            "رتبه کل",
            "نمره",
            "نمره نهائی",
        ]
    )
    sheet.append(
        [
            "الف",
            "ب",
            "0012345690",
            "کلاس ۱",
            "ریاضی نهم",
            5,
            5,
            0,
            0,
            100,
            100,
            1,
            6000,
            1,
            score,
            score,
        ]
    )
    workbook.save(path)
    workbook.close()


def test_discovery_keeps_subject_exam_files_separate_from_comprehensive_files():
    discovered = discover_subject_exam_workbooks(DATA_EXCEL)

    assert {path.name for path in discovered} == set(SUBJECT_EXAM_FILENAMES)
    assert not any(path.name == "701 702 (2).xlsx" for path in discovered)


@pytest.mark.parametrize("filename,expected_rows", SUBJECT_EXAM_ROW_COUNTS.items())
def test_parser_reads_all_rows_and_keeps_formula_and_cached_values(filename, expected_rows):
    parsed = SubjectExamWorkbookParser(DATA_EXCEL / filename).parse()

    assert len(parsed.rows) == expected_rows
    assert parsed.spec.grade_order in {7, 8, 9}
    first = parsed.rows[0]
    assert first.source_row == 2
    assert first.raw_values["نام کاربری"]
    assert first.raw_values["نمره"].startswith("=")
    assert first.evaluated_fields["score"] is not None
    assert first.raw_fields["score"].startswith("=")


@pytest.mark.django_db
def test_ingest_persists_all_three_files_and_is_idempotent(base_data):
    Student.objects.create(
        organization=base_data["organization"],
        national_id="0980506328",
        first_name="امیررضا",
        last_name="امیری",
        birth_date="2012-01-01",
        gender=Student.Gender.MALE,
    )

    first = ingest_subject_exam_workbooks(base_data["school1"], DATA_EXCEL)
    expected_total = sum(SUBJECT_EXAM_ROW_COUNTS.values())

    assert first["discovered_files"] == list(SUBJECT_EXAM_FILENAMES)
    assert first["missing_files"] == []
    assert first["rows_seen"] == expected_total
    assert first["rows_persisted"] == expected_total
    assert first["rows_created"] == expected_total
    assert first["row_write_errors"] == 0
    assert first["rows_valid"] > 0
    assert first["rows_unmatched"] > 0
    assert first["rows_invalid"] > 0
    assert first["rows_valid"] + first["rows_unmatched"] + first["rows_invalid"] == expected_total
    assert SubjectExamResult.objects.count() == expected_total

    first_row = SubjectExamResult.objects.get(
        source_file="کارنامه تفصیلی آزمون تابستانه نهم.xlsx", source_row=2
    )
    assert first_row.status == SubjectExamResult.Status.VALID
    assert first_row.student is not None
    assert first_row.raw_values["نمره"].startswith("=")
    assert str(first_row.score) == "16.0000"

    second = ingest_subject_exam_workbooks(base_data["school1"], DATA_EXCEL)

    assert second["rows_seen"] == expected_total
    assert second["rows_persisted"] == expected_total
    assert second["rows_created"] == 0
    assert second["rows_updated"] == expected_total
    assert SubjectExamResult.objects.count() == expected_total


@pytest.mark.django_db
def test_ingest_keeps_invalid_row_raw_and_does_not_roll_back_valid_row(base_data, tmp_path):
    filename = SUBJECT_EXAM_FILENAMES[0]
    path = tmp_path / filename
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(
        [
            "نام",
            "نام خانوادگی",
            "نام کاربری",
            "کلاس",
            "درس",
            "تعداد سوال",
            "صحیح",
            "غلط",
            "سفید",
            "درصد",
            "بالاترین درصد",
            "رتبه",
            "تراز",
            "رتبه کل",
            "نمره",
            "نمره نهائی",
        ]
    )
    sheet.append(
        [
            "دانش‌آموز سالم",
            "آزمون",
            "0012345690",
            "کلاس ۱",
            "ریاضی نهم",
            5,
            5,
            0,
            0,
            100,
            100,
            1,
            6000,
            1,
            20,
            20,
        ]
    )
    sheet.append(
        [
            "دانش‌آموز خراب",
            "آزمون",
            "شناسه-خراب",
            "کلاس ۱",
            "ریاضی نهم",
            5,
            2,
            2,
            1,
            40,
            100,
            2,
            4000,
            2,
            "نمره-خراب",
            8,
        ]
    )
    workbook.save(path)

    Student.objects.create(
        organization=base_data["organization"],
        national_id="0012345690",
        first_name="دانش‌آموز سالم",
        last_name="آزمون",
        birth_date="2012-01-01",
        gender=Student.Gender.MALE,
    )

    result = ingest_subject_exam_workbooks(base_data["school1"], tmp_path)

    assert result["rows_seen"] == 2
    assert result["rows_persisted"] == 2
    assert result["rows_valid"] == 1
    assert result["rows_invalid"] == 1
    assert result["row_write_errors"] == 0
    invalid = SubjectExamResult.objects.get(source_row=3)
    assert invalid.status == SubjectExamResult.Status.INVALID
    assert invalid.national_id == ""
    assert invalid.score is None
    assert invalid.raw_values["نام کاربری"] == "شناسه-خراب"
    assert invalid.raw_values["نمره"] == "نمره-خراب"
    assert {error["code"] for error in invalid.errors} >= {
        "invalid_national_id",
        "invalid_number",
    }


@pytest.mark.django_db
def test_ingest_links_only_matching_manifest_checksum(base_data, tmp_path):
    path = tmp_path / SUBJECT_EXAM_FILENAMES[0]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["نام", "نام خانوادگی", "نام کاربری", "کلاس", "درس"])
    sheet.append(["الف", "ب", "0012345690", "کلاس ۱", "ریاضی نهم"])
    workbook.save(path)
    checksum = _file_checksum(path)
    manifest = DataSourceManifest.objects.create(
        organization=base_data["organization"],
        school=base_data["school1"],
        source_file=path.name,
        checksum=checksum,
        file_size=path.stat().st_size,
    )

    result = ingest_subject_exam_workbooks(base_data["school1"], tmp_path, manifests=[manifest])

    assert result["rows_persisted"] == 1
    row = SubjectExamResult.objects.get(source_file=path.name, source_row=2)
    assert row.source_manifest_id == manifest.id


@pytest.mark.django_db
def test_changed_source_checksum_retires_old_live_rows(base_data, tmp_path):
    path = tmp_path / SUBJECT_EXAM_FILENAMES[0]
    _write_minimal_subject_exam(path, 20)

    first = ingest_subject_exam_workbooks(base_data["school1"], tmp_path)
    assert first["rows_created"] == 1
    assert SubjectExamResult.objects.count() == 1

    _write_minimal_subject_exam(path, 19)
    second = ingest_subject_exam_workbooks(base_data["school1"], tmp_path)

    assert second["rows_retired"] == 1
    assert second["rows_created"] == 1
    assert SubjectExamResult.objects.count() == 1
    assert SubjectExamResult.all_objects.count() == 2
    assert SubjectExamResult.objects.get().score == 19
