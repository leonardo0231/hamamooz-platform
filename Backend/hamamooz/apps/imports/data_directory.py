"""Repository-data workbook discovery with provenance and conflict protection.

This module is intentionally an *intake* layer.  It does not write Students,
Enrollments, scores, or official reports.  A workbook first becomes a manifest and
its non-empty worksheet rows are stored both raw and normalized.  A later,
explicit class-source selection is required before an official report may use it.
"""

from __future__ import annotations

import hashlib
import os
import re
import unicodedata
from collections import defaultdict
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from django.conf import settings
from django.db import transaction
from openpyxl import load_workbook

from .models import ClassSourceSelection, DataSourceConflict, DataSourceManifest, DataSourceRow

# The input workbooks use a mixture of Persian labels, Latin labels, and old/new
# variants of Arabic characters.  Canonical matching keeps the raw label untouched.
HEADER_ALIASES = {
    "national_id": {"کد ملی", "کدملی", "national id", "national_id", "national code"},
    "student_number": {"شماره دانش‌آموزی", "شماره دانش آموزی", "شماره دانشجویی", "student number"},
    "first_name": {"نام", "first name", "first_name"},
    "last_name": {"نام خانوادگی", "نام خانوادگي", "last name", "last_name"},
    "full_name": {"نام و نام خانوادگی", "نام و نام خانوادگي", "نام دانش‌آموز"},
    "class_code": {"کد کلاس", "کلاس", "class code", "class"},
    "class_name": {"نام کلاس", "نام كلاس", "class name"},
    "grade": {"پایه", "پايه", "پایه تحصیلی", "grade"},
    "month_title": {"ماه", "دوره", "نوبت", "month", "period"},
}

# Summer is the starting period in the supplied source workbooks.  This is a
# period sequence (not the Jalali calendar month number): therefore شهریور=3.
MONTH_NUMBERS = {
    "تیر": 1,
    "مرداد": 2,
    "شهریور": 3,
    "مهر": 4,
    "آبان": 5,
    "آذر": 6,
    "دی": 7,
    "بهمن": 8,
    "اسفند": 9,
    "فروردین": 10,
    "اردیبهشت": 11,
    "خرداد": 12,
}


def data_excel_directory(path: str | Path | None = None) -> Path:
    """Resolve Data/Excel independent of whether Django runs from Backend or root."""

    if path:
        candidate = Path(path)
    elif configured := (
        os.getenv("HAMAMOOZ_DATA_DIRECTORY", "") or getattr(settings, "DATA_EXCEL_DIRECTORY", None)
    ):
        candidate = Path(configured)
    else:
        # BASE_DIR is Backend in normal deployments.  Keep the small parent search
        # for management commands invoked from a checkout, an image, or tests.
        base = Path(settings.BASE_DIR).resolve()
        candidates = [base.parent / "Data" / "Excel", base / "Data" / "Excel"]
        candidate = next((item for item in candidates if item.exists()), candidates[0])
    return candidate.expanduser().resolve()


def _canonical(value: Any) -> str:
    value = unicodedata.normalize("NFKC", str(value or ""))
    value = value.replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    value = re.sub(r"\s+", " ", value).strip().lower()
    return value


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    # openpyxl values are normally primitive.  This fallback ensures malformed
    # formula/error values cannot abort an otherwise auditable intake run.
    if value is not None and not isinstance(value, str | int | float | bool | list | dict):
        return str(value)
    return value


def _text(value: Any) -> str:
    return str(value or "").strip()


def _digits(value: Any) -> str:
    translation = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    return _text(value).translate(translation).replace(" ", "")


def _national_id(value: Any) -> str:
    value = _digits(value)
    if not value.isdigit() or value in {"0", "0000000000"}:
        return ""
    # Excel commonly removes the leading zero.  Do not truncate oversized IDs;
    # preserve them as raw data but leave them non-queryable if malformed.
    if len(value) == 9:
        return value.zfill(10)
    return value if len(value) == 10 else ""


def _header_map(values: Iterable[Any]) -> dict[str, int]:
    aliases = {key: {_canonical(alias) for alias in options} for key, options in HEADER_ALIASES.items()}
    result: dict[str, int] = {}
    for index, value in enumerate(values):
        label = _canonical(value)
        for key, options in aliases.items():
            if label in options and key not in result:
                result[key] = index
    return result


def _normalise_row(headers: list[Any], values: list[Any], header_map: dict[str, int]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for field, index in header_map.items():
        if index < len(values):
            normalized[field] = _json_value(values[index])

    national_id = _national_id(normalized.get("national_id"))
    if national_id:
        normalized["national_id"] = national_id
    else:
        # A header row has the literal value "کد ملی".  Do not let that label
        # masquerade as an actual student identity or evaluation record.
        normalized.pop("national_id", None)
    class_code = _text(normalized.get("class_code"))
    if class_code:
        normalized["class_code"] = class_code
    title = _text(normalized.get("month_title"))
    if title:
        normalized["month_title"] = title
        normalized["month_no"] = MONTH_NUMBERS.get(title)

    # Indicator columns retain their semantic code, independent of their position.
    metrics = {}
    for header, value in zip(headers, values, strict=False):
        match = re.match(r"\s*([A-Z]{2,8}_\d{1,3})\b", _text(header))
        if match:
            metrics[match.group(1)] = _json_value(value)
    if metrics:
        normalized["metrics"] = metrics
    return normalized


def _row_kind(normalized: dict[str, Any], sheet_name: str) -> str:
    if normalized.get("national_id") and normalized.get("month_title"):
        return DataSourceRow.RowKind.EVALUATION
    if normalized.get("national_id"):
        return DataSourceRow.RowKind.STUDENT
    if normalized.get("class_code") or "کلاس" in _canonical(sheet_name):
        return DataSourceRow.RowKind.CLASS
    return DataSourceRow.RowKind.OTHER


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class DataDirectoryScanner:
    """Scan every XLSX file into provenance tables without generating reports."""

    def __init__(self, school, directory: str | Path | None = None):
        self.school = school
        self.directory = data_excel_directory(directory)

    def scan(self) -> dict[str, Any]:
        if not self.directory.is_dir():
            raise FileNotFoundError(f"پوشه داده پیدا نشد: {self.directory}")

        manifests = []
        for path in sorted(self.directory.glob("*.xlsx"), key=lambda item: item.name.casefold()):
            manifests.append(self._scan_workbook(path))
        conflicts = self._refresh_conflicts()
        return {
            "directory": str(self.directory),
            "files_scanned": len(manifests),
            "manifests": [
                {
                    "source_file": item.source_file,
                    "status": item.status,
                    "rows": item.row_count,
                    "student_rows": item.student_row_count,
                    "classes": item.detected_classes,
                    "errors": item.errors,
                }
                for item in manifests
            ],
            "conflicts": conflicts,
        }

    def _scan_workbook(self, path: Path) -> DataSourceManifest:
        source_file = path.relative_to(self.directory).as_posix()
        checksum = _digest(path)
        manifest, _ = DataSourceManifest.all_objects.update_or_create(
            school=self.school,
            source_file=source_file,
            defaults={
                "organization": self.school.organization,
                "checksum": checksum,
                "file_size": path.stat().st_size,
                "status": DataSourceManifest.Status.VALID,
                "detected_classes": [],
                "sheet_manifest": [],
                "errors": [],
                "row_count": 0,
                "student_row_count": 0,
                "is_deleted": False,
                "deleted_at": None,
            },
        )
        # A deterministic rescan replaces only the *derived cache* of raw rows for
        # this unchanged source path; it never touches Students or official reports.
        manifest.rows.all().hard_delete()

        errors: list[dict[str, str]] = []
        sheet_manifest: list[dict[str, Any]] = []
        records: list[DataSourceRow] = []
        detected_classes: set[str] = set()
        class_definitions: set[str] = set()
        student_count = 0
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    sheet_records, sheet_info, classes, definitions = self._read_sheet(manifest, sheet)
                    records.extend(sheet_records)
                    sheet_manifest.append(sheet_info)
                    detected_classes.update(classes)
                    class_definitions.update(definitions)
                    student_count += sum(
                        row.row_kind in {DataSourceRow.RowKind.STUDENT, DataSourceRow.RowKind.EVALUATION}
                        for row in sheet_records
                    )
            finally:
                workbook.close()
        except Exception as exc:  # captured in the manifest, not silently skipped
            errors.append({"code": "workbook_unreadable", "message": str(exc)})

        status = DataSourceManifest.Status.VALID
        if errors:
            status = DataSourceManifest.Status.INVALID
        elif not class_definitions:
            # Crucially do not use a number in the filename as a class definition.
            # `902.xlsx` is a known example: it remains available for investigation
            # but cannot yield an official report.
            status = DataSourceManifest.Status.INCOMPLETE
            errors.append(
                {
                    "code": "missing_class_definition",
                    "message": "فایل کلاس‌بندی معتبر ندارد؛ ایجاد کارنامه رسمی مجاز نیست.",
                }
            )

        DataSourceRow.objects.bulk_create(records, batch_size=500)
        manifest.checksum = checksum
        manifest.file_size = path.stat().st_size
        manifest.status = status
        manifest.detected_classes = sorted(detected_classes)
        manifest.sheet_manifest = sheet_manifest
        manifest.errors = errors
        manifest.row_count = len(records)
        manifest.student_row_count = student_count
        manifest.is_deleted = False
        manifest.deleted_at = None
        manifest.save()
        return manifest

    def _read_sheet(self, manifest, sheet):
        rows = list(sheet.iter_rows(values_only=True))
        header_row_number = None
        header_values: list[Any] = []
        header_map: dict[str, int] = {}
        for number, values in enumerate(rows[:30], start=1):
            candidate = _header_map(values)
            if "national_id" in candidate or "class_code" in candidate:
                header_row_number, header_values, header_map = number, list(values), candidate
                break

        is_classification_sheet = "کلاس" in _canonical(sheet.title) and "class_code" in header_map
        sheet_info = {
            "sheet_name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "header_row": header_row_number,
            "headers": [str(value or "") for value in header_values],
        }
        records: list[DataSourceRow] = []
        classes: set[str] = set()
        definitions: set[str] = set()
        for number, values_tuple in enumerate(rows, start=1):
            values = list(values_tuple)
            if not any(value not in (None, "") for value in values):
                continue
            normalized = _normalise_row(header_values, values, header_map) if header_map else {}
            kind = _row_kind(normalized, sheet.title)
            class_code = _text(normalized.get("class_code"))
            if class_code and number != header_row_number:
                classes.add(class_code)
                if is_classification_sheet:
                    definitions.add(class_code)
            raw = {
                "headers": [str(value or "") for value in header_values] if header_values else [],
                "values": [_json_value(value) for value in values],
            }
            records.append(
                DataSourceRow(
                    manifest=manifest,
                    source_file=manifest.source_file,
                    sheet_name=sheet.title,
                    source_row=number,
                    row_kind=kind,
                    raw_values=raw,
                    normalized_data=normalized,
                    national_id=_national_id(normalized.get("national_id")),
                    class_code=class_code,
                    month_no=normalized.get("month_no"),
                )
            )
        return records, sheet_info, classes, definitions

    @transaction.atomic
    def _refresh_conflicts(self) -> dict[str, int]:
        manifests = list(
            DataSourceManifest.objects.filter(school=self.school).exclude(
                status=DataSourceManifest.Status.INVALID
            )
        )
        # Conflicts are a recalculated view.  Resolved, human-authored records are
        # intentionally retained as audit evidence.
        DataSourceConflict.objects.filter(
            school=self.school, status=DataSourceConflict.Status.OPEN, details__scanner_managed=True
        ).hard_delete()

        by_checksum: dict[str, list[DataSourceManifest]] = defaultdict(list)
        by_class: dict[str, set[DataSourceManifest]] = defaultdict(set)
        by_student: dict[tuple[str, str], set[DataSourceManifest]] = defaultdict(set)
        for manifest in manifests:
            by_checksum[manifest.checksum].append(manifest)
            for class_code in manifest.detected_classes:
                by_class[class_code].add(manifest)
        for row in DataSourceRow.objects.filter(manifest__in=manifests).exclude(national_id=""):
            by_student[(row.class_code, row.national_id)].add(row.manifest)

        counts = defaultdict(int)

        def create(kind, group, manifests_for_conflict, class_code="", national_id=""):
            if len(manifests_for_conflict) < 2:
                return
            conflict = DataSourceConflict.objects.create(
                school=self.school,
                conflict_type=kind,
                class_code=class_code,
                national_id=national_id,
                details={
                    "scanner_managed": True,
                    "group": group,
                    "source_files": sorted(item.source_file for item in manifests_for_conflict),
                    "message": "منبع اصلی باید به‌صورت دستی انتخاب شود.",
                },
            )
            conflict.manifests.set(manifests_for_conflict)
            counts[kind] += 1

        for checksum, candidates in by_checksum.items():
            create(DataSourceConflict.ConflictType.DUPLICATE_FILE, checksum, candidates)
        for class_code, candidates in by_class.items():
            create(DataSourceConflict.ConflictType.CLASS_OVERLAP, class_code, candidates, class_code=class_code)
        for (class_code, national_id), candidates in by_student.items():
            create(
                DataSourceConflict.ConflictType.STUDENT_OVERLAP,
                f"{class_code}:{national_id}",
                candidates,
                class_code=class_code,
                national_id=national_id,
            )

        conflict_classes = {
            item.class_code
            for item in DataSourceConflict.objects.filter(
                school=self.school,
                status=DataSourceConflict.Status.OPEN,
                conflict_type__in=[
                    DataSourceConflict.ConflictType.CLASS_OVERLAP,
                    DataSourceConflict.ConflictType.STUDENT_OVERLAP,
                ],
            ).exclude(class_code="")
        }
        if conflict_classes:
            # JSONField ``overlap`` is PostgreSQL-specific.  Keep this scanner
            # usable with the SQLite development/test setup as well.
            for manifest in manifests:
                if (
                    manifest.status == DataSourceManifest.Status.VALID
                    and conflict_classes.intersection(manifest.detected_classes)
                ):
                    manifest.status = DataSourceManifest.Status.CONFLICT
                    manifest.save(update_fields=["status", "updated_at"])
        return dict(counts)


def selected_manifest_for_class(school, class_code: str) -> DataSourceManifest | None:
    """Return the manually selected valid source, never an implicit last file."""

    selection = (
        ClassSourceSelection.objects.select_related("manifest")
        .filter(school=school, class_code=str(class_code).strip())
        .first()
    )
    if not selection or not selection.manifest:
        return None
    if selection.manifest.status in {
        DataSourceManifest.Status.INCOMPLETE,
        DataSourceManifest.Status.INVALID,
    }:
        return None
    return selection.manifest
