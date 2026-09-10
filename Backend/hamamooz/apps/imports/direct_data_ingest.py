"""Direct, tolerant ingestion of the repository ``Data`` directory.

The files in ``Data/Excel`` are registrar exports, not user-uploaded import
templates.  This service therefore has a different contract from
``ImportJob``: every workbook is inventoried first, every raw row remains
available through :class:`DataSourceRow`, and a malformed row does not discard
the valid rows next to it.

The twelve comprehensive workbooks are written to the existing student,
enrollment and monthly-evaluation tables.  The three summer subject workbooks
are deliberately delegated to the separate subject-exam ingestor so they can
never be confused with monthly indicators.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from hamamooz.apps.evaluations.catalog import FRAMEWORK_VERSION, metric_catalog_for
from hamamooz.apps.evaluations.models import MetricScore, MonthlyEvaluation
from hamamooz.apps.organizations.models import AcademicYear, ClassSection, GradeLevel
from hamamooz.apps.students.models import Enrollment, Student

from .comprehensive import CLASS_SHEET, EVALUATION_SHEET, STUDENT_SHEET
from .comprehensive_flexible import normalize_birth_date
from .data_directory import DataDirectoryScanner, data_excel_directory
from .models import DataSourceManifest, DataSourceRow
from .services.photo_importer import StudentPhotoImporter

COMPREHENSIVE_SHEETS = {CLASS_SHEET, STUDENT_SHEET, EVALUATION_SHEET}
MONTH_NUMBERS = {
    "تیر": 1,
    "مرداد": 2,
    "شهریور": 3,
    "مهر": 4,
    "آبان": 5,
    "آذر": 6,
    "دی": 7,
    "بهمن": 8,
    "اسفند": 9,
    "فروردین": 10,
    "اردیبهشت": 11,
    "خرداد": 12,
}
GRADE_WORDS = {
    "اول": 1,
    "دوم": 2,
    "سوم": 3,
    "چهارم": 4,
    "پنجم": 5,
    "ششم": 6,
    "هفتم": 7,
    "هشتم": 8,
    "نهم": 9,
    "دهم": 10,
    "یازدهم": 11,
    "دوازدهم": 12,
}
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().translate(PERSIAN_DIGITS)


def _canonical(value: Any) -> str:
    value = unicodedata.normalize("NFKC", _text(value))
    value = value.replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    return re.sub(r"\s+", " ", value).strip().lower()


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if value is not None and not isinstance(value, str | int | float | bool | list | dict):
        return str(value)
    return value


def _identifier(value: Any) -> str:
    digits = "".join(character for character in _text(value) if character.isdigit())
    if not digits or digits in {"0", "0000000000"} or len(digits) > 10:
        return ""
    return digits.zfill(10)


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _restore(instance) -> None:
    if instance.is_deleted:
        instance.is_deleted = False
        instance.deleted_at = None


def _month_number(value: Any) -> int | None:
    raw = _canonical(value)
    if raw in {_canonical(key) for key in MONTH_NUMBERS}:
        return next(number for title, number in MONTH_NUMBERS.items() if _canonical(title) == raw)
    try:
        number = int(Decimal(_text(value)))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return number if 1 <= number <= 12 else None


def _grade_order(value: Any) -> int | None:
    raw = _canonical(value).replace("پایه", "").strip()
    if raw.isdigit():
        number = int(raw)
        return number if 1 <= number <= 12 else None
    return GRADE_WORDS.get(raw)


def _class_db_code(value: Any) -> str:
    """Return a deterministic ASCII key while leaving the source code raw.

    Some ninth-grade exports use values such as ``نهم3`` as a class code.
    ``ClassSection.code`` is a slug field, so the technical key must be safe for
    that column; the original value is still retained in every source row and
    in the class title.
    """

    raw = _text(value)
    if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]*", raw):
        return raw[:30]
    match = re.fullmatch(
        r"(اول|دوم|سوم|چهارم|پنجم|ششم|هفتم|هشتم|نهم|دهم|یازدهم|دوازدهم)\s*([A-Za-z0-9۰-۹]+)", raw
    )
    if match:
        grade = _grade_order(match.group(1))
        suffix = _text(match.group(2))
        return f"grade-{grade}-{suffix}"[:30]
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]
    return f"class-{digest}"


def _gender(value: Any) -> str | None:
    raw = _canonical(value)
    if any(token in raw for token in ("پسر", "مذکر", "male", "مرد")):
        return Student.Gender.MALE
    if any(token in raw for token in ("دختر", "مونث", "female", "زن")):
        return Student.Gender.FEMALE
    return None


def _is_comprehensive(path: Path) -> bool:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            names = {_canonical(sheet.title) for sheet in workbook.worksheets}
            return {_canonical(item) for item in COMPREHENSIVE_SHEETS}.issubset(names)
        finally:
            workbook.close()
    except Exception:
        return False


@dataclass
class _ClassRow:
    source_row: int
    school_code: Any
    academic_year_code: Any
    code: Any
    title: Any
    grade: Any
    capacity: Any


@dataclass
class _StudentRow:
    source_row: int
    local_code: Any
    national_id: Any
    student_number: Any
    first_name: Any
    last_name: Any
    gender: Any
    birth_date: Any
    class_code: Any


@dataclass
class _EvaluationRow:
    source_row: int
    sequence: Any
    month: Any
    local_code: Any
    national_id: Any
    full_name: Any
    class_code: Any
    metrics: dict[str, Any]
    note: Any


@dataclass
class _WorkbookRows:
    classes: list[_ClassRow] = field(default_factory=list)
    students: list[_StudentRow] = field(default_factory=list)
    evaluations: list[_EvaluationRow] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class _IngestState:
    students: dict[str, Student] = field(default_factory=dict)
    classes: dict[tuple[str, str], ClassSection] = field(default_factory=dict)
    enrollments: dict[tuple[str, str], Enrollment] = field(default_factory=dict)
    local_students: dict[tuple[str, str], Student] = field(default_factory=dict)


def _read_comprehensive(path: Path) -> _WorkbookRows:
    result = _WorkbookRows()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {_canonical(sheet.title): sheet for sheet in workbook.worksheets}
        classes_sheet = sheets.get(_canonical(CLASS_SHEET))
        students_sheet = sheets.get(_canonical(STUDENT_SHEET))
        evaluation_sheet = sheets.get(_canonical(EVALUATION_SHEET))
        if not classes_sheet or not students_sheet or not evaluation_sheet:
            result.errors.append({"code": "missing_comprehensive_sheet"})
            return result

        for number, values_tuple in enumerate(
            classes_sheet.iter_rows(min_row=5, values_only=True), 5
        ):
            values = list(values_tuple)
            if len(values) < 7 or not any(value not in (None, "") for value in values[1:7]):
                continue
            result.classes.append(
                _ClassRow(number, values[1], values[2], values[3], values[4], values[5], values[6])
            )

        for number, values_tuple in enumerate(
            students_sheet.iter_rows(min_row=5, values_only=True), 5
        ):
            values = list(values_tuple)
            if len(values) < 9 or not any(value not in (None, "") for value in values[2:9]):
                continue
            result.students.append(
                _StudentRow(
                    number,
                    values[1],
                    values[2],
                    values[3],
                    values[4],
                    values[5],
                    values[6],
                    values[7],
                    values[8],
                )
            )

        header_row = None
        headers: list[Any] = []
        for number, values_tuple in enumerate(
            evaluation_sheet.iter_rows(max_row=30, values_only=True), 1
        ):
            values = list(values_tuple)
            if any(re.match(r"\s*[A-Z]{2,8}_\d{1,3}\b", _text(value)) for value in values):
                header_row, headers = number, values
                break
        if header_row is None:
            result.errors.append({"code": "missing_evaluation_header"})
            return result
        metric_columns = {
            index: match.group(1)
            for index, value in enumerate(headers)
            if (match := re.match(r"\s*([A-Z]{2,8}_\d{1,3})\b", _text(value)))
        }
        for number, values_tuple in enumerate(
            evaluation_sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
        ):
            values = list(values_tuple)
            metrics = {
                code: values[index]
                for index, code in metric_columns.items()
                if index < len(values) and values[index] not in (None, "")
            }
            note = values[93] if len(values) > 93 else None
            if not metrics and note in (None, ""):
                continue
            result.evaluations.append(
                _EvaluationRow(
                    source_row=number,
                    sequence=values[0] if len(values) > 0 else None,
                    month=values[1] if len(values) > 1 else None,
                    local_code=values[2] if len(values) > 2 else None,
                    national_id=values[3] if len(values) > 3 else None,
                    full_name=values[4] if len(values) > 4 else None,
                    class_code=values[5] if len(values) > 5 else None,
                    metrics=metrics,
                    note=note,
                )
            )
    finally:
        workbook.close()
    return result


class DirectDataIngestor:
    """Ingest all mounted repository data without an interactive ImportJob."""

    def __init__(self, school, *, excel_directory=None, photo_directory=None, recorded_by=None):
        self.school = school
        self.organization = school.organization
        self.excel_directory = data_excel_directory(excel_directory)
        self.photo_directory = (
            Path(photo_directory).expanduser().resolve()
            if photo_directory
            else self.excel_directory.parent / "Photo"
        )
        self.recorded_by = recorded_by or self._default_actor()
        self.state = _IngestState()
        self.issues: list[dict[str, Any]] = []

    @staticmethod
    def _default_actor():
        from hamamooz.apps.accounts.models import User

        return (
            User.objects.filter(is_superuser=True, is_active=True)
            .order_by("date_joined", "id")
            .first()
            or User.objects.filter(is_staff=True, is_active=True)
            .order_by("date_joined", "id")
            .first()
        )

    def _issue(self, source_file, source_row, code, message, **extra):
        item = {
            "source_file": source_file,
            "source_row": source_row,
            "code": code,
            "message": message,
        }
        item.update(extra)
        self.issues.append(item)
        return item

    def _academic_year(self, raw_code: Any) -> AcademicYear | None:
        key = _canonical(raw_code)
        years = list(AcademicYear.objects.filter(organization=self.organization))
        match = next(
            (
                year
                for year in years
                if key and key in {_canonical(year.code), _canonical(year.title)}
            ),
            None,
        )
        if match:
            return match
        current = [year for year in years if year.is_current and year.is_active]
        active = [year for year in years if year.is_active]
        return (
            current[0]
            if len(current) == 1
            else active[0]
            if len(active) == 1
            else years[0]
            if len(years) == 1
            else None
        )

    def _grade(self, raw_value: Any) -> GradeLevel | None:
        order = _grade_order(raw_value)
        grades = GradeLevel.objects.filter(organization=self.organization)
        if order is not None:
            match = grades.filter(order=order).first()
            if match:
                return match
        key = _canonical(raw_value)
        return next((grade for grade in grades if key == _canonical(grade.title)), None)

    def _upsert_class(self, source_file: str, row: _ClassRow) -> ClassSection | None:
        code = _text(row.code)
        if not code or code == "کد کلاس":
            return None
        db_code = _class_db_code(code)
        year = self._academic_year(row.academic_year_code)
        grade = self._grade(row.grade)
        if year is None or grade is None:
            self._issue(
                source_file,
                row.source_row,
                "class_reference",
                "سال تحصیلی یا پایه کلاس قابل تطبیق نیست.",
            )
            return None
        try:
            capacity = int(Decimal(_text(row.capacity))) if _text(row.capacity) else 30
        except (InvalidOperation, ValueError):
            self._issue(
                source_file,
                row.source_row,
                "class_capacity",
                "ظرفیت کلاس عددی نیست؛ مقدار خام حفظ شد.",
            )
            capacity = 30
        capacity = max(1, min(capacity, 32767))
        key = (str(year.id), code)
        existing = (
            self.state.classes.get(key)
            or ClassSection.all_objects.filter(
                school=self.school, academic_year=year, code=db_code
            ).first()
        )
        if existing:
            if (
                existing.grade_level_id != grade.id
                or (row.title and existing.title != _text(row.title))
                or existing.capacity != capacity
            ):
                self._issue(
                    source_file,
                    row.source_row,
                    "class_conflict",
                    "کلاس قبلاً از منبع دیگری با مشخصات متفاوت ثبت شده؛ منبع اول به‌صورت قطعی نگه داشته شد.",
                    existing_title=existing.title,
                    incoming_title=_text(row.title),
                )
            _restore(existing)
            existing.is_active = True
            existing.save(update_fields=["is_deleted", "deleted_at", "is_active", "updated_at"])
            self.state.classes[key] = existing
            return existing
        title = _text(row.title) or code
        section = ClassSection(
            school=self.school,
            academic_year=year,
            grade_level=grade,
            code=db_code,
            title=title[:100],
            capacity=capacity,
        )
        try:
            section.full_clean(exclude=["id"])
            section.save()
        except Exception as exc:
            self._issue(source_file, row.source_row, "class_write", str(exc))
            return None
        self.state.classes[key] = section
        return section

    def _upsert_student(self, source_file: str, row: _StudentRow) -> Student | None:
        national_id = _identifier(row.national_id)
        if not national_id:
            self._issue(
                source_file,
                row.source_row,
                "student_identity",
                "کد ملی معتبر نیست؛ ردیف خام حفظ شد.",
                raw_id=_text(row.national_id),
            )
            return None
        first_name, last_name = _text(row.first_name), _text(row.last_name)
        gender = _gender(row.gender)
        try:
            birth_date, _ = normalize_birth_date(row.birth_date)
        except ValueError as exc:
            birth_date = None
            self._issue(
                source_file,
                row.source_row,
                "student_birth_date",
                str(exc),
                raw_date=_text(row.birth_date),
            )
        existing = (
            self.state.students.get(national_id)
            or Student.all_objects.filter(
                organization=self.organization, national_id=national_id
            ).first()
        )
        if existing:
            _restore(existing)
            if birth_date is None or gender is None or not first_name or not last_name:
                self._issue(
                    source_file,
                    row.source_row,
                    "student_partial",
                    "مشخصات ردیف کامل نیست؛ مشخصات معتبر قبلی حفظ شد.",
                )
            else:
                incoming = (first_name, last_name, birth_date, gender)
                current = (
                    existing.first_name,
                    existing.last_name,
                    existing.birth_date,
                    existing.gender,
                )
                if current != incoming:
                    self._issue(
                        source_file,
                        row.source_row,
                        "student_conflict",
                        "مشخصات دانش‌آموز با رکورد قبلی متفاوت است؛ منبع اول حفظ شد.",
                    )
            existing.status = Student.Status.ACTIVE
            existing.save(update_fields=["is_deleted", "deleted_at", "status", "updated_at"])
            self.state.students[national_id] = existing
            return existing
        if not first_name or not last_name or birth_date is None or gender is None:
            self._issue(
                source_file,
                row.source_row,
                "student_write",
                "برای ایجاد دانش‌آموز نام، نام خانوادگی، تاریخ تولد و جنسیت معتبر لازم است.",
            )
            return None
        student = Student(
            organization=self.organization,
            national_id=national_id,
            first_name=first_name[:100],
            last_name=last_name[:100],
            birth_date=birth_date,
            gender=gender,
            status=Student.Status.ACTIVE,
        )
        try:
            student.full_clean(exclude=["id"])
            student.save()
        except Exception as exc:
            self._issue(source_file, row.source_row, "student_write", str(exc))
            return None
        self.state.students[national_id] = student
        return student

    def _upsert_enrollment(
        self, source_file: str, row: _StudentRow, student: Student
    ) -> Enrollment | None:
        code = _text(row.class_code)
        section = next(
            (
                item
                for (year_id, class_code), item in self.state.classes.items()
                if class_code == code
            ),
            None,
        )
        if section is None:
            self._issue(
                source_file,
                row.source_row,
                "enrollment_class",
                "کلاس دانش‌آموز در شیت کلاس‌بندی ثبت نشده است.",
                class_code=code,
            )
            return None
        year = section.academic_year
        key = (str(student.id), str(year.id))
        existing = (
            self.state.enrollments.get(key)
            or Enrollment.all_objects.filter(student=student, academic_year=year)
            .order_by("-updated_at")
            .first()
        )
        student_number = (_text(row.student_number) or student.national_id)[:50]
        if existing:
            if (
                existing.status == Enrollment.Status.ACTIVE
                and existing.class_section_id != section.id
            ):
                self._issue(
                    source_file,
                    row.source_row,
                    "enrollment_conflict",
                    "دانش‌آموز در همین سال قبلاً در کلاس دیگری ثبت شده؛ رکورد اول حفظ شد.",
                )
                self.state.enrollments[key] = existing
                return existing
            _restore(existing)
            existing.school = self.school
            existing.grade_level = section.grade_level
            existing.class_section = section
            existing.student_number = student_number
            existing.status = Enrollment.Status.ACTIVE
            existing.enrolled_on = year.starts_on
            existing.left_on = None
            try:
                existing.full_clean(exclude=["id"])
                existing.save()
            except Exception as exc:
                self._issue(source_file, row.source_row, "enrollment_write", str(exc))
                return None
            self.state.enrollments[key] = existing
            return existing
        duplicate_number = Enrollment.objects.filter(
            school=self.school,
            academic_year=year,
            student_number=student_number,
            status=Enrollment.Status.ACTIVE,
        ).first()
        if duplicate_number:
            self._issue(
                source_file,
                row.source_row,
                "enrollment_number_conflict",
                "شماره دانش‌آموزی قبلاً برای دانش‌آموز دیگری ثبت شده است.",
            )
            return None
        enrollment = Enrollment(
            student=student,
            school=self.school,
            academic_year=year,
            grade_level=section.grade_level,
            class_section=section,
            student_number=student_number,
            status=Enrollment.Status.ACTIVE,
            enrolled_on=year.starts_on,
        )
        try:
            enrollment.full_clean(exclude=["id"])
            enrollment.save()
        except Exception as exc:
            self._issue(source_file, row.source_row, "enrollment_write", str(exc))
            return None
        self.state.enrollments[key] = enrollment
        return enrollment

    def _metric_score(
        self, value: Any, source_file: str, source_row: int, metric_code: str
    ) -> int | None:
        try:
            decimal = Decimal(_text(value))
        except (InvalidOperation, ValueError):
            self._issue(
                source_file,
                source_row,
                "metric_value",
                f"امتیاز {metric_code} عددی نیست؛ مقدار خام حفظ شد.",
            )
            return None
        integer = int(decimal)
        if decimal == integer and 0 <= integer <= 5:
            return integer
        if decimal == integer and 5 < integer <= 20 and integer % 4 == 0:
            return integer // 4
        self._issue(
            source_file,
            source_row,
            "metric_value",
            f"امتیاز {metric_code} در مقیاس قابل ثبت ۰ تا ۵ نیست؛ مقدار خام حفظ شد.",
            raw_value=_text(value),
        )
        return None

    def _upsert_evaluation(
        self, source_file: str, row: _EvaluationRow, enrollment: Enrollment
    ) -> int:
        month_no = _month_number(row.month)
        if month_no is None:
            self._issue(source_file, row.source_row, "evaluation_month", "ماه ارزیابی معتبر نیست.")
            return 0
        if self.recorded_by is None:
            self._issue(
                source_file,
                row.source_row,
                "evaluation_actor",
                "کاربر ثبت‌کننده پیدا نشد؛ داده خام حفظ شد.",
            )
            return 0
        valid_metrics = {}
        catalog = metric_catalog_for(FRAMEWORK_VERSION) or {}
        for metric_code, value in row.metrics.items():
            if catalog and metric_code not in catalog:
                self._issue(
                    source_file,
                    row.source_row,
                    "metric_code",
                    f"کد شاخص {metric_code} در چارچوب فعال نیست.",
                )
                continue
            score = self._metric_score(value, source_file, row.source_row, metric_code)
            if score is not None:
                valid_metrics[metric_code] = score
        evaluation = MonthlyEvaluation.all_objects.filter(
            enrollment=enrollment, month_no=month_no, framework_version=FRAMEWORK_VERSION
        ).first()
        if evaluation is None:
            evaluation = MonthlyEvaluation(
                enrollment=enrollment,
                month_no=month_no,
                framework_version=FRAMEWORK_VERSION,
                note=_text(row.note)[:5000],
                recorded_by=self.recorded_by,
            )
        else:
            _restore(evaluation)
            evaluation.note = _text(row.note)[:5000]
            evaluation.recorded_by = self.recorded_by
        try:
            evaluation.save()
        except Exception as exc:
            self._issue(source_file, row.source_row, "evaluation_write", str(exc))
            return 0
        if not valid_metrics:
            return 0
        try:
            now = timezone.now()
            MetricScore.objects.bulk_create(
                [
                    MetricScore(
                        evaluation=evaluation,
                        metric_code=metric_code,
                        value=score,
                        created_at=now,
                        updated_at=now,
                    )
                    for metric_code, score in valid_metrics.items()
                ],
                update_conflicts=True,
                update_fields=["value", "updated_at"],
                unique_fields=["evaluation", "metric_code"],
            )
            return len(valid_metrics)
        except Exception as exc:
            # Keep the tolerant contract even on a backend that does not support
            # conflict-aware bulk inserts; fall back to isolated writes.
            written = 0
            for metric_code, score in valid_metrics.items():
                try:
                    MetricScore.objects.update_or_create(
                        evaluation=evaluation, metric_code=metric_code, defaults={"value": score}
                    )
                    written += 1
                except Exception as metric_exc:
                    self._issue(
                        source_file,
                        row.source_row,
                        "metric_write",
                        str(metric_exc),
                        metric_code=metric_code,
                    )
            if written == 0:
                self._issue(source_file, row.source_row, "metric_bulk_write", str(exc))
            return written

    def _mark_source_row(
        self, manifest, sheet_name, source_row, *, enrollment=None, status="linked", **details
    ):
        source = DataSourceRow.objects.filter(
            manifest=manifest, sheet_name=sheet_name, source_row=source_row
        ).first()
        if not source:
            return
        normalized = dict(source.normalized_data or {})
        normalized["direct_ingest"] = {"status": status, **details}
        source.normalized_data = normalized
        if enrollment is not None:
            source.enrollment = enrollment
        source.save(update_fields=["normalized_data", "enrollment", "updated_at"])

    @staticmethod
    def _mark_manifest(manifest, summary, *, status):
        """Persist the write result separately from the scanner's shape status."""

        manifest.ingest_status = status
        manifest.ingest_summary = summary
        manifest.ingested_at = timezone.now()
        manifest.save(
            update_fields=["ingest_status", "ingest_summary", "ingested_at", "updated_at"]
        )

    def _ingest_comprehensive(self, path: Path, manifest: DataSourceManifest) -> dict[str, Any]:
        source_file = path.relative_to(self.excel_directory).as_posix()
        issue_start = len(self.issues)
        summary = {
            "file": source_file,
            "students_seen": 0,
            "students_written": 0,
            "enrollments_written": 0,
            "evaluations_seen": 0,
            "evaluations_written": 0,
            "metric_scores_written": 0,
            "errors": 0,
        }
        rows = _read_comprehensive(path)
        for error in rows.errors:
            self._issue(
                source_file, None, error.get("code", "workbook"), "ساختار فایل جامع کامل نیست."
            )
        local_enrollments: dict[str, Enrollment] = {}
        for class_row in rows.classes:
            self._upsert_class(source_file, class_row)
            self._mark_source_row(
                manifest, CLASS_SHEET, class_row.source_row, status="class_processed"
            )
        for student_row in rows.students:
            summary["students_seen"] += 1
            student = self._upsert_student(source_file, student_row)
            if student is None:
                summary["errors"] += 1
                self._mark_source_row(
                    manifest, STUDENT_SHEET, student_row.source_row, status="invalid"
                )
                continue
            summary["students_written"] += 1
            enrollment = self._upsert_enrollment(source_file, student_row, student)
            if enrollment:
                summary["enrollments_written"] += 1
                local_enrollments[_text(student_row.local_code)] = enrollment
                self._mark_source_row(
                    manifest,
                    STUDENT_SHEET,
                    student_row.source_row,
                    enrollment=enrollment,
                    status="linked",
                )
            else:
                self._mark_source_row(
                    manifest, STUDENT_SHEET, student_row.source_row, status="student_only"
                )
        for evaluation_row in rows.evaluations:
            summary["evaluations_seen"] += 1
            national_id = _identifier(evaluation_row.national_id)
            student = self.state.students.get(national_id) if national_id else None
            enrollment = None
            if student:
                candidates = [
                    item
                    for (student_id, _year_id), item in self.state.enrollments.items()
                    if student_id == str(student.id)
                ]
                if evaluation_row.class_code:
                    candidates = [
                        item
                        for item in candidates
                        if item.class_section.code == _text(evaluation_row.class_code)
                    ] or candidates
                enrollment = candidates[0] if candidates else None
            if enrollment is None and _text(evaluation_row.local_code):
                enrollment = local_enrollments.get(_text(evaluation_row.local_code))
            if enrollment is None:
                self._issue(
                    source_file,
                    evaluation_row.source_row,
                    "evaluation_identity",
                    "ثبت‌نام متناظر برای ارزیابی پیدا نشد.",
                )
                summary["errors"] += 1
                self._mark_source_row(
                    manifest, EVALUATION_SHEET, evaluation_row.source_row, status="unmatched"
                )
                continue
            before = MonthlyEvaluation.objects.filter(
                enrollment=enrollment,
                month_no=_month_number(evaluation_row.month),
                framework_version=FRAMEWORK_VERSION,
            ).exists()
            metric_count = self._upsert_evaluation(source_file, evaluation_row, enrollment)
            after = MonthlyEvaluation.objects.filter(
                enrollment=enrollment,
                month_no=_month_number(evaluation_row.month),
                framework_version=FRAMEWORK_VERSION,
            ).exists()
            if after:
                summary["evaluations_written"] += 0 if before else 1
                summary["metric_scores_written"] += metric_count or 0
                self._mark_source_row(
                    manifest,
                    EVALUATION_SHEET,
                    evaluation_row.source_row,
                    enrollment=enrollment,
                    status="linked",
                )
            else:
                summary["errors"] += 1
        summary["errors"] = len(self.issues) - issue_start
        return summary

    def _ingest_subject_files(
        self, files: list[Path], manifests: dict[str, DataSourceManifest]
    ) -> dict[str, Any]:
        if not files:
            # Keep the shape compatible with the real subject-exam ingestor so
            # the manifest update loop can remain uniform when a deployment
            # intentionally mounts only the comprehensive workbooks.
            return {
                "files": [],
                "rows_seen": 0,
                "rows_persisted": 0,
                "rows_created": 0,
                "rows_updated": 0,
                "rows_valid": 0,
                "rows_invalid": 0,
                "rows_unmatched": 0,
                "row_write_errors": 0,
                "errors": [{"code": "no_subject_workbooks"}],
            }
        from .services.subject_exam import ingest_subject_exam_workbooks

        return ingest_subject_exam_workbooks(
            self.school,
            self.excel_directory,
            manifests=manifests,
        )

    def _ingest_photos(self) -> dict[str, Any]:
        if not self.photo_directory.is_dir():
            return {"directory": str(self.photo_directory), "status": "missing"}
        importer = StudentPhotoImporter(self.organization)
        result = {
            "directory": str(self.photo_directory),
            "directory_result": importer.import_directory(self.photo_directory),
        }
        archives = []
        for path in sorted(self.photo_directory.rglob("*.zip"), key=lambda item: str(item)):
            try:
                archives.append(
                    {
                        "file": str(path.relative_to(self.photo_directory)),
                        "result": importer.import_zip(path),
                    }
                )
            except ValueError as exc:
                archives.append(
                    {"file": str(path.relative_to(self.photo_directory)), "error": str(exc)}
                )
        result["archives"] = archives
        result["unsupported_archives"] = [
            str(path.relative_to(self.photo_directory))
            for path in sorted(self.photo_directory.rglob("*.rar"), key=lambda item: str(item))
            if path.is_file()
        ]
        return result

    def run(self) -> dict[str, Any]:
        if not self.excel_directory.is_dir():
            raise FileNotFoundError(f"پوشه داده پیدا نشد: {self.excel_directory}")
        scan = DataDirectoryScanner(self.school, self.excel_directory).scan()
        manifests = {
            manifest.source_file: manifest
            for manifest in DataSourceManifest.objects.filter(school=self.school)
        }
        files = sorted(self.excel_directory.glob("*.xlsx"), key=lambda item: item.name.casefold())
        comprehensive_files = [path for path in files if _is_comprehensive(path)]
        subject_files = [path for path in files if path not in comprehensive_files]
        comprehensive_results = []
        for path in comprehensive_files:
            source_file = path.relative_to(self.excel_directory).as_posix()
            manifest = manifests.get(source_file)
            if manifest is None:
                continue
            try:
                with transaction.atomic():
                    result = self._ingest_comprehensive(path, manifest)
                    comprehensive_results.append(result)
                    self._mark_manifest(
                        manifest,
                        result,
                        status=(
                            DataSourceManifest.IngestStatus.COMPLETED
                            if result.get("errors", 0) == 0
                            else DataSourceManifest.IngestStatus.PARTIAL
                        ),
                    )
            except Exception as exc:
                self._issue(source_file, None, "file_write", str(exc))
                result = {"file": source_file, "errors": 1, "fatal_error": str(exc)}
                comprehensive_results.append(result)
                self._mark_manifest(manifest, result, status=DataSourceManifest.IngestStatus.FAILED)
        subject_result = self._ingest_subject_files(subject_files, manifests)
        for item in subject_result.get("files", []):
            manifest = manifests.get(item.get("source_file"))
            if manifest is None:
                continue
            self._mark_manifest(
                manifest,
                item,
                status=(
                    DataSourceManifest.IngestStatus.COMPLETED
                    if item.get("row_write_errors", 0) == 0 and not item.get("errors")
                    else DataSourceManifest.IngestStatus.PARTIAL
                ),
            )
        photo_result = self._ingest_photos()
        return {
            "phase_order": ["comprehensive", "subject_exam", "photos"],
            "directory": str(self.excel_directory),
            "files_scanned": scan.get("files_scanned", 0),
            "comprehensive_files": len(comprehensive_files),
            "subject_files": len(subject_files),
            "comprehensive": comprehensive_results,
            "subject_exam": subject_result,
            "photos": photo_result,
            "issues": self.issues[:2000],
            "issue_count": len(self.issues),
            "completed_at": timezone.now().isoformat(),
        }


def ingest_data_directory(school, **kwargs) -> dict[str, Any]:
    """Convenience entry point used by the management command and Docker."""

    return DirectDataIngestor(school, **kwargs).run()
